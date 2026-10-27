"""Documents: PyMuPDF (PDF/EPUB/XPS), pandoc (text formats), Edge (HTML -> PDF) and LibreOffice when installed."""
import glob
import os
import pathlib
import re
import shutil
import subprocess
import uuid
import zipfile

from . import config

try:
    import fitz
except ImportError:
    fitz = None
try:
    import pypandoc
    PANDOC = pypandoc.get_pandoc_path()
except (ImportError, OSError):
    PANDOC = shutil.which("pandoc")

CREATE_NO_WINDOW = 0x08000000
HERE = os.path.dirname(os.path.abspath(__file__))
CSS = os.path.join(HERE, "doc.css")


def edge_exe():
    for p in (os.path.expandvars(r"%ProgramFiles(x86)%\Microsoft\Edge\Application\msedge.exe"),
              os.path.expandvars(r"%ProgramFiles%\Microsoft\Edge\Application\msedge.exe")):
        if os.path.exists(p):
            return p
    return None


def soffice_exe():
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return shutil.which("soffice")


def _run(ctx, cmd, cwd=None, timeout=900):
    proc = subprocess.Popen(cmd, cwd=cwd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            creationflags=CREATE_NO_WINDOW)
    ctx.attach(proc)
    try:
        out, err = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        proc.kill()
        raise RuntimeError("Timed out")
    if ctx.cancelled:
        raise InterruptedError
    if proc.returncode != 0:
        msg = (err or out or b"").decode("utf-8", "replace").strip().splitlines()
        raise RuntimeError((msg[-1] if msg else "failed")[:300])
    return out


def _zip_dir(folder, dst_zip):
    with zipfile.ZipFile(dst_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(folder):
            for f in files:
                full = os.path.join(root, f)
                z.write(full, os.path.relpath(full, folder))
    return dst_zip


def _pages(spec, n):
    """'1-3,5' -> [0,1,2,4] (clipped to the document)."""
    if not spec or not str(spec).strip():
        return list(range(n))
    out = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        a, _, b = part.partition("-")
        try:
            lo = int(a) if a else 1
            hi = int(b) if b else (n if _ else lo)
        except ValueError:
            continue
        out += [i - 1 for i in range(max(lo, 1), min(hi, n) + 1)]
    return sorted(set(out)) or list(range(n))


# ---------------------------------------------------------------- PyMuPDF
def pdf_facts(path):
    try:
        doc = fitz.open(path)
    except Exception:
        return {}
    info = {"pages": doc.page_count, "encrypted": doc.needs_pass}
    if doc.needs_pass:
        return info
    try:
        p = doc[0]
        info["page_size"] = f"{p.rect.width / 72 * 25.4:.0f}x{p.rect.height / 72 * 25.4:.0f} mm"
        sample = range(min(doc.page_count, 5))
        chars = sum(len(doc[i].get_text().strip()) for i in sample)
        imgs = sum(len(doc[i].get_images()) for i in sample)
        info["text"] = chars > 40
        info["scanned"] = chars < 40 and imgs > 0
        info["images"] = imgs
        meta = doc.metadata or {}
        if meta.get("title"):
            info["title"] = meta["title"][:80]
    except Exception:
        pass
    return info


def fitz_thumbnail(path, out):
    doc = fitz.open(path)
    page = doc[0]
    zoom = 480 / max(page.rect.width, page.rect.height, 1)
    page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).save(out)


def fitz_convert(ctx, src, dst, sfid, dfid):
    """PDF/EPUB/XPS/MOBI/FB2/CBZ -> pdf, txt, html, images, svg; PDF -> optimised PDF."""
    doc = fitz.open(src)
    if doc.needs_pass:
        raise RuntimeError("This PDF is password protected")
    q = ctx.opts.get("quality") or "balanced"
    if dfid == "pdf" and sfid == "pdf":
        return _pdf_optimize(ctx, doc, dst, q)
    if dfid == "pdf":
        ctx.progress(None, "Laying out pages")
        pdf = fitz.open("pdf", doc.convert_to_pdf())
        pdf.save(dst, garbage=3, deflate=True)
        return dst
    pages = _pages(ctx.opts.get("pages"), doc.page_count)
    if dfid == "txt":
        with open(dst, "w", encoding="utf-8") as f:
            for i, n in enumerate(pages):
                f.write(doc[n].get_text().rstrip() + "\n\n")
                ctx.progress((i + 1) / len(pages), f"Page {n + 1}")
        return dst
    if dfid == "html":
        parts = []
        for i, n in enumerate(pages):
            parts.append(f'<section class="page">{doc[n].get_text("xhtml")}</section>')
            ctx.progress((i + 1) / len(pages), f"Page {n + 1}")
        title = os.path.splitext(os.path.basename(src))[0]
        with open(CSS, encoding="utf-8") as f:
            css = f.read()
        with open(dst, "w", encoding="utf-8") as f:
            f.write(f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title>"
                    f"<style>{css}</style></head><body>{''.join(parts)}</body></html>")
        return dst
    if dfid in ("png", "jpg", "webp", "tiff", "svg"):
        dpi = {"original": 300, "high": 200, "balanced": 150, "small": 96}[q]
        stem = os.path.splitext(os.path.basename(dst))[0]
        folder = os.path.join(os.path.dirname(dst), stem + "_pages")
        os.makedirs(folder, exist_ok=True)
        files = []
        for i, n in enumerate(pages):
            name = os.path.join(folder, f"{stem}-{n + 1:03d}.{dfid}")
            page = doc[n]
            if dfid == "svg":
                with open(name, "w", encoding="utf-8") as f:
                    f.write(page.get_svg_image())
            else:
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                if dfid in ("png",):
                    pix.save(name)
                else:
                    from PIL import Image
                    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    kw = {"jpg": dict(quality=88, optimize=True), "webp": dict(quality=85),
                          "tiff": dict(compression="tiff_lzw", dpi=(dpi, dpi))}[dfid]
                    im.save(name, {"jpg": "JPEG", "webp": "WEBP", "tiff": "TIFF"}[dfid], **kw)
            files.append(name)
            if ctx.cancelled:
                raise InterruptedError
            ctx.progress((i + 1) / len(pages), f"Page {n + 1} of {doc.page_count}")
        if len(files) == 1:
            shutil.move(files[0], dst)
            return dst
        return _zip_dir(folder, os.path.splitext(dst)[0] + ".zip")
    raise ValueError(f"cannot make {dfid}")


def _pdf_optimize(ctx, doc, dst, q):
    ctx.progress(0.2, "Recompressing images")
    target = {"original": None, "high": (220, 85), "balanced": (150, 75), "small": (100, 55)}[q]
    if target and hasattr(doc, "rewrite_images"):
        dpi, quality = target
        try:
            doc.rewrite_images(dpi_threshold=dpi + 30, dpi_target=dpi, quality=quality, lossy=True, lossless=True,
                               bitonal=True, color=True, gray=True)
        except Exception:
            pass
    ctx.progress(0.7, "Cleaning up")
    try:
        doc.subset_fonts()
    except Exception:
        pass
    doc.save(dst, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True, use_objstms=1)
    return dst


def pdf_to_docx(ctx, src, dst, sfid, dfid):
    from pdf2docx import Converter
    ctx.progress(None, "Rebuilding layout (this can take a while)")
    doc = fitz.open(src)
    pages = _pages(ctx.opts.get("pages"), doc.page_count)
    doc.close()
    cv = Converter(src)
    try:
        cv.convert(dst, pages=pages)
    finally:
        cv.close()
    return dst


def merge_pdf(ctx, items, dst):
    """Merge PDFs and pictures (in the given order) into one PDF."""
    from . import images
    out = fitz.open()
    for i, (path, fid) in enumerate(items):
        if fid == "pdf":
            src = fitz.open(path)
        elif fid in ("epub", "xps", "mobi", "fb2", "cbz", "svg"):
            src = fitz.open("pdf", fitz.open(path).convert_to_pdf())
        else:
            src = fitz.open("pdf", images.to_pdf_bytes(path, fid))
        out.insert_pdf(src)
        ctx.progress((i + 1) / len(items), f"Added {os.path.basename(path)}")
        if ctx.cancelled:
            raise InterruptedError
    out.save(dst, garbage=3, deflate=True)
    return dst


# ---------------------------------------------------------------- pandoc
READ = {"md": "markdown", "txt": "markdown+hard_line_breaks", "html": "html", "docx": "docx", "odt": "odt",
        "rtf": "rtf", "epub": "epub", "tex": "latex", "rst": "rst", "org": "org", "textile": "textile",
        "wiki": "mediawiki", "ipynb": "ipynb", "fb2": "fb2", "typ": "typst", "pptx": "pptx", "adoc": "asciidoc"}
WRITE = {"md": "gfm", "txt": "plain", "html": "html5", "docx": "docx", "odt": "odt", "rtf": "rtf", "epub": "epub3",
         "tex": "latex", "rst": "rst", "org": "org", "textile": "textile", "wiki": "mediawiki", "ipynb": "ipynb",
         "fb2": "fb2", "typ": "typst", "pptx": "pptx", "adoc": "asciidoc"}
MEDIA_OUT = {"md", "rst", "org", "textile", "wiki", "typ", "adoc", "tex"}   # text formats that link images


def pandoc_convert(ctx, src, dst, sfid, dfid):
    if not PANDOC:
        raise RuntimeError("pandoc is missing (pip install pypandoc_binary)")
    work = os.path.dirname(dst)
    title = os.path.splitext(os.path.basename(src))[0]
    # --sandbox: pandoc reads only the files named here, so a received document can't pull other local files
    # (or web images) into the result
    cmd = [PANDOC, src, "-f", READ[sfid], "-t", WRITE[dfid], "-o", dst, "--wrap=preserve", "--sandbox"]
    if dfid in ("html", "rtf", "tex", "typ", "adoc"):
        cmd += ["--standalone"]
    if dfid == "html":
        cmd += ["--embed-resources", "--css", CSS, "--metadata", f"pagetitle={title}"]
    if dfid == "epub":
        cmd += ["--metadata", f"title={title}"]
    if dfid in MEDIA_OUT:
        cmd += ["--extract-media=media"]
    if dfid == "docx" and sfid in ("md", "txt", "html"):
        cmd += ["--toc-depth=3"]
    ctx.progress(None, "Converting with pandoc")
    try:
        _run(ctx, cmd, cwd=work)
    except RuntimeError:
        if dfid != "html":
            raise
        _run(ctx, [c for c in cmd if c != "--embed-resources"], cwd=work)  # e.g. unreachable web images
    media = os.path.join(work, "media")
    if dfid in MEDIA_OUT and os.path.isdir(media) and os.listdir(media):
        bundle = os.path.join(work, "bundle")
        os.makedirs(bundle, exist_ok=True)
        shutil.move(dst, os.path.join(bundle, os.path.basename(dst)))
        shutil.move(media, os.path.join(bundle, "media"))
        return _zip_dir(bundle, os.path.splitext(dst)[0] + ".zip")
    return dst


# ---------------------------------------------------------------- Edge: HTML -> PDF
def html_to_pdf(ctx, src, dst, sfid, dfid):
    edge = edge_exe()
    if not edge:
        raise RuntimeError("Microsoft Edge is needed for PDF output")
    ctx.progress(None, "Printing to PDF")
    profile = os.path.join(config.DATA_DIR, "edge-print")
    if os.path.exists(dst):
        os.remove(dst)
    # The page is untrusted. Edge prints a copy that starts with a strict Content-Security-Policy (no scripts, nothing
    # from the web; local pictures and styles next to the page still work through <base>), and every request goes to
    # an unreachable proxy, loopback included. (--blink-settings=scriptEnabled=false would stop printing altogether.)
    safe = _locked_copy(src, os.path.dirname(dst))
    try:
        _run(ctx, [edge, "--headless=new", "--disable-gpu", "--no-first-run", "--disable-extensions",
                   f"--user-data-dir={profile}", "--no-pdf-header-footer", "--virtual-time-budget=4000",
                   "--proxy-server=http://127.0.0.1:9", "--proxy-bypass-list=<-loopback>",
                   "--disable-background-networking", "--disable-sync", "--disable-component-update",
                   f"--print-to-pdf={dst}", pathlib.Path(safe).as_uri()], timeout=180)
    finally:
        try:
            os.remove(safe)
        except OSError:
            pass
    if not os.path.exists(dst):
        raise RuntimeError("Edge did not produce a PDF")
    return dst


PRINT_CSP = ("default-src 'none'; script-src 'none'; object-src 'none'; frame-src 'none'; connect-src 'none'; "
             "img-src file: data:; media-src file: data:; font-src file: data:; style-src 'unsafe-inline' file: data:")


def _locked_copy(src, folder):
    """A copy of the HTML that Edge may print: CSP + <base> to the original folder (percent-encoded file URL)."""
    with open(src, "rb") as f:
        data = f.read()
    bom = b""
    for b in (b"\xef\xbb\xbf",):
        if data.startswith(b):
            bom, data = b, data[len(b):]
    base = pathlib.Path(os.path.dirname(os.path.abspath(src))).as_uri() + "/"
    head = (f'<meta http-equiv="Content-Security-Policy" content="{PRINT_CSP}">'
            f'<base href="{base}">').encode("ascii", "xmlcharrefreplace")
    out = os.path.join(folder, f"print-{uuid.uuid4().hex[:8]}.html")
    with open(out, "wb") as f:
        f.write(bom + head + data)
    return out


# ---------------------------------------------------------------- LibreOffice (optional)
LO_FILTER = {"pdf": "pdf", "docx": "docx:MS Word 2007 XML", "doc": "doc:MS Word 97", "odt": "odt", "rtf": "rtf",
             "xlsx": "xlsx:Calc MS Excel 2007 XML", "xls": "xls:MS Excel 97", "ods": "ods", "csv": "csv",
             "pptx": "pptx:Impress MS PowerPoint 2007 XML", "ppt": "ppt:MS PowerPoint 97", "odp": "odp",
             "html": "html", "txt": "txt:Text (encoded):UTF8"}


def office_convert(ctx, src, dst, sfid, dfid):
    exe = soffice_exe()
    if not exe:
        raise RuntimeError("LibreOffice is not installed")
    ctx.progress(None, "Converting with LibreOffice")
    work = os.path.dirname(dst)
    outdir = os.path.join(work, "lo-out")
    os.makedirs(outdir, exist_ok=True)
    profile = "file:///" + os.path.join(config.DATA_DIR, "lo-profile").replace("\\", "/")
    _run(ctx, [exe, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--convert-to",
               LO_FILTER[dfid], "--outdir", outdir, src], timeout=600)
    made = glob.glob(os.path.join(outdir, "*"))
    if not made:
        raise RuntimeError("LibreOffice produced no file")
    shutil.move(made[0], dst)
    return dst


def text_to_html(ctx, src, dst, sfid, dfid):
    """Plain text kept exactly as written (line breaks, spacing), ready for PDF printing."""
    import html
    raw = open(src, "rb").read()
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = raw.decode(enc).replace("\r\n", "\n")
            break
        except UnicodeDecodeError:
            continue
    with open(CSS, encoding="utf-8") as f:
        css = f.read()
    with open(dst, "w", encoding="utf-8") as f:
        f.write(f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head>"
                f"<body><pre class='plain'>{html.escape(text)}</pre></body></html>")
    return dst


def word_count(path, fid):
    if fid == "docx":
        try:
            import docx
            d = docx.Document(path)
            words = sum(len(p.text.split()) for p in d.paragraphs)
            return {"words": words, "images": len(d.inline_shapes), "tables": len(d.tables)}
        except Exception:
            return {}
    if fid not in ("txt", "md", "html", "rst", "org", "tex", "adoc", "textile", "wiki", "typ"):
        return {}
    try:
        with open(path, "rb") as f:
            text = f.read(5_000_000).decode("utf-8", "ignore")
    except OSError:
        return {}
    if fid == "html":
        text = re.sub(r"<[^>]+>", " ", text)
    return {"words": len(text.split()), "lines": text.count("\n") + 1}
