"""Tables (pandas), structured data (JSON/YAML/TOML/XML), archives and fonts."""
import csv
import html
import json
import os
import shutil
import sqlite3
import tarfile
import zipfile


# ---------------------------------------------------------------- tables
def _sniff_sep(path):
    with open(path, "rb") as f:
        sample = f.read(65536).decode("utf-8-sig", "ignore")
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def read_tables(path, fid):
    """{sheet name: DataFrame}"""
    import pandas as pd
    if fid in ("csv", "tsv"):
        sep = "\t" if fid == "tsv" else _sniff_sep(path)
        for enc in ("utf-8-sig", "cp1250", "latin-1"):
            try:
                return {"data": pd.read_csv(path, sep=sep, encoding=enc, low_memory=False)}
            except UnicodeDecodeError:
                continue
    if fid in ("xlsx", "xls", "ods"):
        return pd.read_excel(path, sheet_name=None, engine={"ods": "odf", "xls": "xlrd"}.get(fid))
    if fid == "parquet":
        return {"data": pd.read_parquet(path)}
    if fid == "xml":
        return {"data": pd.read_xml(path)}
    if fid == "sqlite":
        con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        try:
            names = [r[0] for r in con.execute("select name from sqlite_master where type='table' "
                                               "and name not like 'sqlite_%'")]
            return {n: pd.read_sql_query(f'select * from "{n}"', con) for n in names}
        finally:
            con.close()
    if fid in ("json", "yaml", "toml"):
        data = read_struct(path, fid)
        if isinstance(data, dict):
            lists = {k: v for k, v in data.items() if isinstance(v, list) and v and isinstance(v[0], dict)}
            if lists:
                return {k: pd.json_normalize(v) for k, v in lists.items()}
            return {"data": pd.json_normalize(data)}
        if isinstance(data, list):
            return {"data": pd.json_normalize(data) if data and isinstance(data[0], dict) else pd.DataFrame({"value": data})}
        raise ValueError("This JSON is not a table")
    raise ValueError(f"cannot read {fid} as a table")


def table_facts(path, fid):
    try:
        if fid in ("xlsx",):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            sheets = wb.sheetnames
            ws = wb[sheets[0]]
            return {"sheets": len(sheets), "rows": ws.max_row, "cols": ws.max_column}
        if fid in ("csv", "tsv"):
            sep = "\t" if fid == "tsv" else _sniff_sep(path)
            with open(path, "rb") as f:
                rows = sum(1 for _ in f)
            with open(path, encoding="utf-8-sig", errors="ignore") as f:
                cols = len(f.readline().split(sep))
            return {"rows": rows, "cols": cols, "sep": {",": "comma", ";": "semicolon", "\t": "tab", "|": "pipe"}[sep]}
    except Exception:
        pass
    return {}


def write_tables(ctx, tables, dst, dfid):
    import pandas as pd
    single = dfid in ("csv", "tsv", "parquet", "json", "xml", "md", "html")
    if single and len(tables) > 1 and dfid not in ("html", "md", "json"):
        # one file per sheet, zipped
        folder = os.path.splitext(dst)[0] + "_sheets"
        os.makedirs(folder, exist_ok=True)
        for i, (name, df) in enumerate(tables.items()):
            safe = "".join(c for c in str(name) if c not in '\\/:*?"<>|') or f"sheet{i + 1}"
            write_tables(ctx, {name: df}, os.path.join(folder, f"{safe}.{dfid}"), dfid)
        out = os.path.splitext(dst)[0] + ".zip"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for f in os.listdir(folder):
                z.write(os.path.join(folder, f), f)
        return out
    ctx.progress(0.6, "Writing")
    if dfid in ("sqlite", "xlsx", "ods", "csv", "tsv", "parquet", "xml"):
        # nested lists/objects (from JSON) become JSON text in a cell
        tables = {k: df.map(lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v)
                  for k, df in tables.items()}
    first = next(iter(tables.values()))
    if dfid == "csv":
        first.to_csv(dst, index=False, encoding="utf-8-sig")
    elif dfid == "tsv":
        first.to_csv(dst, index=False, sep="\t", encoding="utf-8")
    elif dfid == "parquet":
        first.columns = [str(c) for c in first.columns]
        first.to_parquet(dst, index=False)
    elif dfid == "xml":
        df = first.copy()
        df.columns = [_xml_name(c) for c in df.columns]
        df.to_xml(dst, index=False, root_name="rows", row_name="row")
    elif dfid == "json":
        data = {k: json.loads(v.to_json(orient="records", date_format="iso", force_ascii=False))
                for k, v in tables.items()}
        with open(dst, "w", encoding="utf-8") as f:
            json.dump(data if len(data) > 1 else next(iter(data.values())), f, ensure_ascii=False, indent=2)
    elif dfid in ("xlsx", "ods"):
        with pd.ExcelWriter(dst, engine="openpyxl" if dfid == "xlsx" else "odf") as w:
            for name, df in tables.items():
                sheet = str(name)[:31] or "Sheet1"
                df.to_excel(w, sheet_name=sheet, index=False)
                if dfid == "xlsx":  # readable column widths + frozen header
                    ws = w.sheets[sheet]
                    ws.freeze_panes = "A2"
                    for i, col in enumerate(df.columns):
                        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
                        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = min(max(width + 2, 6), 60)
    elif dfid == "sqlite":
        if os.path.exists(dst):
            os.remove(dst)
        con = sqlite3.connect(dst)
        try:
            for name, df in tables.items():
                df.to_sql(str(name), con, index=False)
        finally:
            con.close()
    elif dfid == "md":
        with open(dst, "w", encoding="utf-8") as f:
            for name, df in tables.items():
                if len(tables) > 1:
                    f.write(f"## {name}\n\n")
                f.write(df.to_markdown(index=False) + "\n\n")
    elif dfid == "html":
        from .docs import CSS
        with open(CSS, encoding="utf-8") as f:
            css = f.read()
        parts = [(f"<h2>{html.escape(str(name))}</h2>" if len(tables) > 1 else "") + df.to_html(index=False, na_rep="", border=0)
                 for name, df in tables.items()]
        with open(dst, "w", encoding="utf-8") as f:
            f.write(f"<!doctype html><html><head><meta charset='utf-8'><style>{css}"
                    f"@page{{size:A4 landscape}} body{{max-width:none}}</style></head><body>{''.join(parts)}</body></html>")
    else:
        raise ValueError(f"cannot write {dfid}")
    return dst


def _xml_name(c):
    s = "".join(ch if ch.isalnum() or ch in "_-." else "_" for ch in str(c))
    return s if s and (s[0].isalpha() or s[0] == "_") else "_" + s


def table_convert(ctx, src, dst, sfid, dfid):
    ctx.progress(0.2, "Reading table")
    tables = read_tables(src, sfid)
    return write_tables(ctx, tables, dst, dfid)


# ---------------------------------------------------------------- structured data
def read_struct(path, fid):
    with open(path, "rb") as f:
        raw = f.read()
    text = raw.decode("utf-8-sig")
    if fid == "json":
        return json.loads(text)
    if fid == "yaml":
        import yaml
        return yaml.safe_load(text)
    if fid == "toml":
        import tomllib
        return tomllib.loads(text)
    if fid == "xml":
        import xmltodict
        return xmltodict.parse(text)
    raise ValueError(fid)


def struct_facts(path, fid):
    if os.path.getsize(path) > 50_000_000:
        return {}
    try:
        d = read_struct(path, fid)
    except Exception as e:
        return {"invalid": str(e)[:120]}
    if isinstance(d, list):
        tab = bool(d) and all(isinstance(x, dict) for x in d[:50])
        return {"items": len(d), "tabular": tab}
    if isinstance(d, dict):
        return {"keys": len(d), "tabular": any(isinstance(v, list) and v and isinstance(v[0], dict) for v in d.values())}
    return {}


def _clean_toml(v):
    if isinstance(v, dict):
        return {str(k): _clean_toml(x) for k, x in v.items() if x is not None}
    if isinstance(v, list):
        return [_clean_toml(x) for x in v if x is not None]
    return v


def struct_convert(ctx, src, dst, sfid, dfid):
    data = read_struct(src, sfid)
    ctx.progress(0.5, "Writing")
    with open(dst, "w", encoding="utf-8") as f:
        if dfid == "json":
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)
        elif dfid == "yaml":
            import yaml
            yaml.safe_dump(json.loads(json.dumps(data, default=str)), f, allow_unicode=True, sort_keys=False)
        elif dfid == "toml":
            import tomli_w
            if not isinstance(data, dict):
                data = {"items": data}
            f.write(tomli_w.dumps(_clean_toml(json.loads(json.dumps(data, default=str)))))
        elif dfid == "xml":
            import xmltodict
            if not (isinstance(data, dict) and len(data) == 1):
                data = {"root": {"item": data} if isinstance(data, list) else data}
            f.write(xmltodict.unparse(json.loads(json.dumps(data, default=str)), pretty=True))
        else:
            raise ValueError(dfid)
    return dst


# ---------------------------------------------------------------- archives
def seven_zip_exe():
    for p in (r"C:\Program Files\7-Zip\7z.exe", r"C:\Program Files (x86)\7-Zip\7z.exe"):
        if os.path.exists(p):
            return p
    return shutil.which("7z")


TAR_MODE = {"tar": "w", "tgz": "w:gz", "tbz2": "w:bz2", "txz": "w:xz"}
MAX_ENTRIES = 100_000              # archive limits (zip bombs, endless lists)
MAX_UNPACKED = 20 * 2 ** 30
MAX_RATIO = 200                    # unpacked size / archive size


def _inside(folder, name):
    """The extraction target for an archive member, or ValueError if it would land outside the folder."""
    base = os.path.realpath(folder)
    name = name.replace("\\", "/")
    if name.startswith("/") or ":" in name or any(p == ".." for p in name.split("/")):
        raise ValueError(f"unsafe path in archive: {name}")
    target = os.path.realpath(os.path.join(base, name))
    if os.path.commonpath([base, target]) != base:
        raise ValueError(f"unsafe path in archive: {name}")
    return target


def _limits(src, count, unpacked):
    free = shutil.disk_usage(os.path.dirname(os.path.abspath(src))).free
    if count > MAX_ENTRIES:
        raise RuntimeError(f"This archive has {count:,} entries; Morph unpacks at most {MAX_ENTRIES:,}.")
    if unpacked > MAX_UNPACKED or unpacked > free - 2 ** 30:
        raise RuntimeError("This archive unpacks to more than Morph can safely hold on this disk.")
    if unpacked > MAX_RATIO * max(os.path.getsize(src), 1) and unpacked > 100 * 2 ** 20:
        raise RuntimeError("This archive expands suspiciously much (a 'zip bomb'?), so Morph won't unpack it.")


def _extract_tar(src, folder):
    """Member by member: only plain files and folders, every path checked. Links and devices are skipped, so even
    the tarfile filter bypasses fixed after Python 3.12.10 (CVE-2025-4517, CVE-2025-4330 and relatives) can't apply."""
    with tarfile.open(src) as t:
        members = t.getmembers()
        _limits(src, len(members), sum(m.size for m in members if m.isfile()))
        for m in members:
            target = _inside(folder, m.name)
            if m.isdir():
                os.makedirs(target, exist_ok=True)
            elif m.isfile():
                os.makedirs(os.path.dirname(target), exist_ok=True)
                with t.extractfile(m) as fin, open(target, "wb") as fout:
                    shutil.copyfileobj(fin, fout, 1 << 20)
            # symlinks, hard links, devices and FIFOs are left out on purpose


def _extract_zip(src, folder):
    with zipfile.ZipFile(src) as z:
        infos = z.infolist()
        _limits(src, len(infos), sum(i.file_size for i in infos))
        for i in infos:
            target = _inside(folder, i.filename)
            if i.is_dir():
                os.makedirs(target, exist_ok=True)
                continue
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with z.open(i) as fin, open(target, "wb") as fout:
                shutil.copyfileobj(fin, fout, 1 << 20)


def extract(ctx, src, sfid, folder):
    if sfid == "zip":
        _extract_zip(src, folder)
    elif sfid in TAR_MODE:
        _extract_tar(src, folder)
    elif sfid == "7z":
        import py7zr
        with py7zr.SevenZipFile(src) as z:
            lst = z.list()
            _limits(src, len(lst), sum(x.uncompressed or 0 for x in lst))
            for x in lst:
                _inside(folder, x.filename)   # refuse the whole archive if any path would escape
            z.extractall(folder)
    elif sfid == "rar":
        exe = seven_zip_exe()
        if not exe:
            raise RuntimeError("RAR needs 7-Zip installed")
        import subprocess
        r = subprocess.run([exe, "x", "-y", f"-o{folder}", src], capture_output=True, creationflags=0x08000000)
        if r.returncode:
            raise RuntimeError("7-Zip could not open this RAR")
    else:
        raise ValueError(sfid)


def archive_facts(path, fid):
    try:
        if fid == "zip":
            with zipfile.ZipFile(path) as z:
                infos = z.infolist()
                return {"files": sum(not i.is_dir() for i in infos), "unpacked": sum(i.file_size for i in infos)}
        if fid in TAR_MODE:
            with tarfile.open(path) as t:
                m = t.getmembers()
                return {"files": sum(x.isfile() for x in m), "unpacked": sum(x.size for x in m)}
        if fid == "7z":
            import py7zr
            with py7zr.SevenZipFile(path) as z:
                lst = z.list()
                return {"files": sum(not x.is_directory for x in lst), "unpacked": sum(x.uncompressed or 0 for x in lst)}
    except Exception:
        pass
    return {}


def archive_convert(ctx, src, dst, sfid, dfid):
    folder = os.path.join(os.path.dirname(dst), "unpacked")
    os.makedirs(folder, exist_ok=True)
    ctx.progress(0.1, "Unpacking")
    extract(ctx, src, sfid, folder)
    files = []
    for root, _, names in os.walk(folder):
        files += [os.path.join(root, n) for n in names]
    q = ctx.opts.get("quality") or "balanced"
    level = {"original": 6, "high": 6, "balanced": 6, "small": 9}[q]
    total = max(len(files), 1)
    if dfid == "zip":
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED, compresslevel=level) as z:
            for i, f in enumerate(files):
                z.write(f, os.path.relpath(f, folder))
                ctx.progress(0.2 + 0.8 * (i + 1) / total, f"Packing {i + 1}/{total}")
                if ctx.cancelled:
                    raise InterruptedError
    elif dfid == "7z":
        import py7zr
        ctx.progress(None, "Compressing (7z)")
        filters = [{"id": py7zr.FILTER_LZMA2, "preset": 9 if q == "small" else 7}]
        with py7zr.SevenZipFile(dst, "w", filters=filters) as z:
            z.writeall(folder, arcname="")
    elif dfid in TAR_MODE:
        kw = {} if dfid in ("tar", "txz") else {"compresslevel": level}
        if dfid == "txz":
            kw = {"preset": 9 if q == "small" else 6}
        with tarfile.open(dst, TAR_MODE[dfid], **kw) as t:
            for i, f in enumerate(files):
                t.add(f, os.path.relpath(f, folder))
                ctx.progress(0.2 + 0.8 * (i + 1) / total, f"Packing {i + 1}/{total}")
                if ctx.cancelled:
                    raise InterruptedError
    else:
        raise ValueError(dfid)
    return dst


# ---------------------------------------------------------------- fonts
def font_facts(path):
    try:
        from fontTools.ttLib import TTFont
        f = TTFont(path, lazy=True)
        name = f["name"].getBestFullName()
        return {"family": name, "outlines": "PostScript (CFF)" if ("CFF " in f or "CFF2" in f) else "TrueType",
                "glyphs": f["maxp"].numGlyphs}
    except Exception:
        return {}


def font_convert(ctx, src, dst, sfid, dfid):
    from fontTools.ttLib import TTFont
    font = TTFont(src)
    cff = "CFF " in font or "CFF2" in font
    if dfid == "ttf" and cff:
        raise RuntimeError("This font has PostScript outlines - pick OTF instead of TTF")
    if dfid == "otf" and not cff:
        raise RuntimeError("This font has TrueType outlines - pick TTF instead of OTF")
    font.flavor = {"woff": "woff", "woff2": "woff2"}.get(dfid)
    ctx.progress(0.5, "Compressing" if font.flavor else "Writing")
    font.save(dst)
    return dst
